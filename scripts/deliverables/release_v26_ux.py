# -*- coding: utf-8 -*-
"""v2.6.0: эргономика в документах. Руководство пользователя — глава 22 (шапка, меню, «Мой день», первые шаги, мастер стрима,
внешний вид) + ссылки в главах 1 и 2; описания — раздел «Мой день» и абзац о навигации; версия во всех .docx.
Снимки — capture_special_shots.py (s_myday, s_hubs, s_menus, s_umenu, s_onb, s_wizard). Идемпотентно.
Запуск: python scripts/deliverables/release_v26_ux.py"""
import glob, os, pathlib, docx
from docx.shared import Inches

os.chdir(pathlib.Path(__file__).resolve().parents[2])
SH = pathlib.Path("build/shots")
P = "Управление_стримами_и_командами"
OLDV, NEWV = "2.5.0", "2.6.0"
CH = "22. Навигация, «Мой день» и быстрые действия"
SECTIONS = [
 ("22.1 Шапка", "В шапке — только то, что нужно постоянно: меню, переключатель «Стрим › Команда», поиск, «＋ Создать», индикатор состояния и меню пользователя.",
  ["«＋ Создать» — быстрое добавление задачи, бага, риска или импедимента команды, фичи или риска стрима (ROAM) и мастер «Новый стрим» — из любого раздела, без перехода.",
   "Индикатор состояния — одна точка: зелёная — синхронизировано, жёлтая — нужен вход или есть проблема, серая — локальный режим; по клику — подробности.",
   "Меню пользователя: роль и вход, данные (Excel, отчёт PPTX, JSON, импорт, сброс), тема, плотность, стартовая страница, видеообзор."],
  [("s_menus.jpg", "Меню «＋ Создать»"), ("s_umenu.jpg", "Меню пользователя")]),
 ("22.2 Меню", "Меню разделено на два уровня — «Стрим» и «Команда»; родственные разделы собраны в один пункт и переключаются вкладками над разделом.",
  ["Уровень «Стрим»: дашборд стрима, фичи, PI, синхронизация (Scrum of Scrums, System Demo), риски ROAM, зрелость, оргструктура (стримы и команды, портфель).",
   "Уровень «Команда»: Мой день, обзор команды, бэклог, приоритизация (RICE, WSJF, MoSCoW), доски, планирование, расписание, церемонии, выпуск, риски и журналы, качество, люди, здоровье, настройки, справка.",
   "Поиск в меню ищет по всем разделам, в том числе внутри пунктов; ⭐ — избранные разделы; «Все разделы» открывает прежнее полное меню из 65 разделов.",
   "Открытый раздел строится сразу, остальные — при открытии, поэтому переключение команд и разделов быстрое."],
  [("s_hubs.jpg", "Компактное меню и вкладки «Приоритизации»")]),
 ("22.3 «Мой день»", "Стартовая страница по умолчанию (меняется в меню пользователя).",
  ["«Я —»: выберите себя в составе команды один раз — дальше показываются ваши задачи текущего спринта; «вся команда» — задачи всех.",
   "Задачи сгруппированы по статусу; «▶ В работу» и «✓ Готово» меняют статус в один клик (с датой начала и завершения).",
   "«Требует внимания»: ваши импедименты и риски, критичные баги, заблокированные зависимости; «Ближайшие 14 дней»: конец спринта, демо, праздники, отпуска, дни рождения.",
   "Сигналы стрима — первые пять с переходом в дашборд стрима (если в стриме больше одной команды)."],
  [("s_myday.jpg", "Стартовая страница «Мой день»")]),
 ("22.4 Первые шаги новой команды", "Пустая команда видит на «Моём дне» и дашборде чек-лист: параметры, состав, бэклог, план спринта, цель спринта, первая сессия на доске. Шаги отмечаются автоматически, клик открывает нужный раздел; «Скрыть» убирает чек-лист.",
  [], [("s_onb.jpg", "Первые шаги команды на дашборде")]),
 ("22.5 Мастер «Новый стрим»", "Стрим, его команды (по одной в строке, пустые или копии текущей) и первый PI с датой старта и числом итераций создаются за один шаг; после создания открывается дашборд нового стрима.",
  [], [("s_wizard.jpg", "Мастер «Новый стрим»")]),
 ("22.6 Внешний вид", "Интерфейс стал спокойнее и плотнее.",
  ["Цвет — только у отклонений: норма нейтральна, подсвечиваются жёлтые и красные значения.",
   "Таблицы выглядят как таблицы: стрелки списков и счётчики чисел появляются при наведении и редактировании.",
   "Описание раздела — одна строка; клик по ней раскрывает текст.",
   "«Плотность: компактно» в меню пользователя уменьшает отступы, шрифт KPI и строк таблиц — больше данных на экране."], []),
]
MYDAY = ("Мой день", "Стартовая страница: ваши задачи текущего спринта со сменой статуса в один клик, блокеры и риски, события ближайших 14 дней "
         "(конец спринта, демо, праздники, отпуска, дни рождения) и сигналы стрима. Для пустой команды здесь же — чек-лист первых шагов.")
NAVTXT = ("Навигация. С версии 2.6 меню разделено на уровни «Стрим» и «Команда», родственные разделы собраны в пункты с вкладками "
          "(например, «Приоритизация» — RICE, WSJF и MoSCoW), а «＋ Создать» в шапке добавляет задачи, баги, риски и фичи из любого раздела.")


def put_before(anchor, p):
    anchor._p.addprevious(p._p); return p


def add_pic(d, anchor, img, cap, width):
    if not (SH / img).exists(): print("   ! нет снимка", img); return
    p = d.add_paragraph(); p.alignment = 1; p.add_run().add_picture(str(SH / img), width=Inches(width)); put_before(anchor, p)
    c = d.add_paragraph(cap); c.alignment = 1
    for r in c.runs: r.italic = True
    put_before(anchor, c)


def version(d):
    pars = list(d.paragraphs) + [p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs] \
        + [p for s in d.sections for hf in (s.header, s.footer) for p in hf.paragraphs]
    n = 0
    for p in pars:
        if OLDV in p.text:
            for r in p.runs: r.text = r.text.replace(OLDV, NEWV)
            n += 1
    return n


# руководство пользователя
fn = f"Руководство_пользователя_{P}.docx"; d = docx.Document(fn); ps = d.paragraphs
if not any(p.text.startswith(CH) for p in ps):
    anchor = next(p for p in ps if p.style.name == "Heading 1" and p.text.startswith("Ширина столбцов"))
    put_before(anchor, d.add_paragraph(CH, style="Heading 1"))
    for title, text, bullets, pics in SECTIONS:
        put_before(anchor, d.add_paragraph(title, style="Heading 2")); put_before(anchor, d.add_paragraph(text))
        for b in bullets: put_before(anchor, d.add_paragraph(b, style="List Bullet"))
        for img, cap in pics: add_pic(d, anchor, img, cap, 6.3)
    ps = d.paragraphs
    toc = next((k for k, p in enumerate(ps) if p.text == "21. Уровень стрима" and p.style.name == "Normal"), None)
    if toc is not None: ps[toc + 1].insert_paragraph_before("22. Навигация, «Мой день» и быстрые действия", style=ps[toc].style)
    for head, note in (("1. Начало работы", "Начните со стартовой страницы «Мой день»: ваши задачи, блокеры и события. Меню, шапка и быстрые действия — глава 22."),
                       ("2. Интерфейс", "С версии 2.6 шапка и меню упрощены: уровни «Стрим | Команда», пункты с вкладками, «＋ Создать» и меню пользователя — см. главу 22.")):
        i = next((k for k, p in enumerate(d.paragraphs) if p.style.name == "Heading 1" and p.text.startswith(head)), None)
        if i is not None: d.paragraphs[i + 1].insert_paragraph_before(note)
    st = "добавлена"
else: st = "уже есть"
d.save(fn); print(fn[:58], "глава 22:", st)
# описания
for fn, pics in ((f"Иллюстрированное_описание_{P}.docx", True), (f"Описание_{P}.docx", False)):
    d = docx.Document(fn)
    if not any(p.style.name.startswith("Heading") and p.text.strip() == MYDAY[0] for p in d.paragraphs):
        ps = d.paragraphs; i = next(k for k, p in enumerate(ps) if p.style.name == "Heading 1" and p.text.strip() == "Обзор")
        anchor = ps[i + 1]
        put_before(anchor, d.add_paragraph(MYDAY[0], style="Heading 2")); put_before(anchor, d.add_paragraph(MYDAY[1]))
        if pics: add_pic(d, anchor, "s_myday.jpg", "Мой день", 6.61)
        j = next((k for k, p in enumerate(d.paragraphs) if p.style.name == "Heading 1" and p.text.startswith("Кросс-функциональная команда")), None)
        if j is not None:
            k = next(x for x in range(j + 1, len(d.paragraphs)) if d.paragraphs[x].style.name.startswith("Heading"))
            d.paragraphs[k].insert_paragraph_before(NAVTXT)
            if pics: add_pic(d, d.paragraphs[k], "s_hubs.jpg", "Меню «Стрим | Команда» и вкладки разделов", 6.61)
        st = "добавлен"
    else: st = "уже есть"
    d.save(fn); print(fn[:58], "«Мой день»:", st)
for fn in sorted(glob.glob(f"*_{P}.docx")):
    d = docx.Document(fn); n = version(d); d.save(fn); print(f"{fn[:58]:58s} версия: {n}")
