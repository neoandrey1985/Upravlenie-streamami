# -*- coding: utf-8 -*-
"""Комплект документов «Управления стримами и командами» v2.3.0 (первый выпуск нового продукта).

1) Во всех .docx (кроме автогенерируемых «Архитектуры» и «Тестовой документации»): название продукта, версия,
   внешний порт 8090 (внутренний api:8080 не трогаем), имена БД streams / streams_audit.
2) Новые главы: стримы и команды, уровень стрима (обзор, фичи, PI, ROAM, Scrum of Scrums, System Demo),
   с снимками из capture_stream_shots.py (build/shots/s_*.jpg) в иллюстрированных документах.
Идемпотентно: повторный запуск не дублирует главы.  Запуск: python scripts/deliverables/release_v23_streams.py
"""
import os, re, copy, pathlib, docx
from docx.shared import Inches

os.chdir(pathlib.Path(__file__).resolve().parents[2])
SH = pathlib.Path("build/shots")
P = "Управление_стримами_и_командами"
UG, ILL, DESC, DEP, ADM = (f"Руководство_пользователя_{P}.docx", f"Иллюстрированное_описание_{P}.docx", f"Описание_{P}.docx",
                           f"Инструкция_по_развёртыванию_{P}.docx", f"Руководство_администратора_{P}.docx")
MARK = "Стримы и команды"

# ---------- 1. замены ----------
LITERAL = [
    ("Управление кросс-функциональной командой", "Управление стримами и командами"),
    ("v1.49.0 · 28.10.2026", "v2.3.0 · 29.10.2026"), ("1.49.0", "2.3.0"),
    ("localhost:8080", "localhost:8090"), ("nginx (порт 8080)", "nginx (порт 8090)"), ("только порт 8080 (nginx)", "только порт 8090 (nginx)"),
    ("через nginx (порт 8080)", "через nginx (порт 8090)"), ("web:8080", "web:8090"), ("8080 → 80 (единственный наружу)", "8090 → 80 (единственный наружу)"),
    ("web 8080, api 8081", "web 8090, api 8081"),
    ("planirovanie_data.json", "streams_data.json"), ("planirovanie_audit", "streams_audit"), ("5432/planirovanie", "5432/streams"),
    ("planirovanie.sql", "streams.sql"), ("/opt/planirovanie", "/opt/streams"), ("Planirovanie analytics", "Streams analytics"),
    ("psycopg2://planirovanie:", "psycopg2://streams:"),
]
WORD = re.compile(r"(?<![\w./-])planirovanie(?![\w-])")   # имя БД/пользователя; planirovanie-api-1.0.0.jar не трогаем


def fix_text(s):
    for a, b in LITERAL: s = s.replace(a, b)
    return WORD.sub("streams", s)


def fix_par(p):
    if fix_text(p.text) == p.text: return 0
    for r in p.runs: r.text = fix_text(r.text)
    if fix_text(p.text) != p.text:                     # шаблон разорван между runs — собираем текст в первом run
        full = fix_text(p.text); runs = p.runs
        runs[0].text = full
        for r in runs[1:]: r.text = ""
    return 1


def all_pars(d):
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            for c in row.cells: yield from c.paragraphs
    for s in d.sections:
        for hf in (s.header, s.footer): yield from hf.paragraphs


# ---------- 2. содержание новых глав ----------
ORG = ("Стримы и команды",
       "Оргструктура «Стрим → Команда». Стрим (поток ценности, Agile Release Train) объединяет несколько кросс-функциональных команд, "
       "которые вместе поставляют ценность. У каждой команды свои данные — бэклог, спринты, доска, риски, метрики; на сервере они хранятся "
       "в отдельной схеме базы данных, в браузере — под своим ключом. Текущая команда выбирается переключателем «Стрим › Команда» в шапке.",
       ["Переключатель в шапке: команды сгруппированы по стримам, текущая отмечена; выбор команды открывает её данные в том же разделе.",
        "Раздел «Стримы и команды»: сводка по командам каждого стрима (участники, задачи, открытые), кнопка «Открыть».",
        "Создание команды: пустая (справочники DoD, DoR, календарь, FAQ — из текущей команды, без рабочих элементов) или копия текущей.",
        "«📥 Команда из JSON» — новая команда из JSON-выгрузки, в том числе из «Управления кросс-функциональной командой».",
        "Переименование, перенос команды в другой стрим, удаление команды вместе с данными (первую команду удалить нельзя), удаление пустого стрима.",
        "Роли (с сервером): назначение на стрим действует на все его команды, на команду — только на неё; без назначений действует общая роль."],
       ["s_org.jpg", "s_switch.jpg"], ["Раздел «Стримы и команды»: два стрима, семь команд", "Переключатель «Стрим › Команда» в шапке"])
STREAM = [
    ("Обзор стрима",
     "Сводка по всем командам текущего стрима, посчитанная теми же формулами, что и дашборд команды: velocity, health, прогресс и риск срыва "
     "текущего спринта, say/do, открытые баги, блокеры и предсказуемость PI. Ниже — графики burn-up PI и velocity команд по итерациям, "
     "фичи текущего PI по WSJF с прогрессом.",
     ["KPI стрима: команд, людей, velocity стрима (сумма средних), health, предсказуемость PI, открытые баги, блокеры, высокие риски ROAM.",
      "Таблица команд с RAG-подсветкой; «Открыть» переключает на команду. Команды без доступа помечаются «нет доступа».",
      "Burn-up PI: объём историй фич PI, сделанное нарастающим итогом и идеальный темп; velocity — столбцы по командам и итерациям.",
      "«⟳ Обновить» перечитывает данные всех команд стрима."],
     ["s_stream.jpg", "s_charts.jpg"], ["Обзор стрима: KPI и таблица команд", "Обзор стрима: burn-up PI и velocity команд по итерациям"]),
    ("Фичи стрима",
     "Бэклог фич стрима с приоритетом WSJF = (BV + TC + RR) / размер, статусом (Воронка → Анализ → Бэклог → В работе → Готово), PI и владельцем. "
     "Фича разбивается на истории прямо в бэклогах команд: «＋ История» создаёт историю в выбранной команде и в спринте нужной итерации PI. "
     "Прогресс фичи считается по статусам этих историй.",
     ["Режимы «Таблица» и «Канбан»: в канбане карточки перетаскиваются между колонками статусов или переносятся кнопками ◀ ▶.",
      "На карточке — WSJF, PI, владелец, прогресс и команды, в которых есть истории фичи.",
      "Список историй фичи раскрывается в колонке «Прогресс» (команда, ID, название, статус)."],
     ["s_features.jpg", "s_kanban.jpg"], ["Фичи стрима: WSJF, статусы и прогресс по историям команд", "Канбан фич по статусам"]),
    ("PI стрима",
     "Планирование инкремента (PI): параметры PI (название, старт, число итераций по 14 дней), программная доска «команды × итерации», "
     "цели команд на PI и межкомандные зависимости.",
     ["Программная доска строится автоматически из историй фич: фича появляется в строке команды в итерации спринта истории; готовые — зелёные.",
      "Цели команд: плановая и фактическая бизнес-ценность (BV), признак «обязательная»; предсказуемость PI = факт / план по обязательным целям.",
      "Межкомандные зависимости: кто ждёт, кто делает, фича, к какой итерации нужно и когда запланировано; конфликт сроков или блокировка подсвечиваются красным и на доске (🔗)."],
     ["s_pi.jpg", "s_deps.jpg"], ["PI стрима: программная доска и цели команд", "Межкомандные зависимости с контролем сроков"]),
    ("Риски ROAM",
     "Риски уровня стрима по модели ROAM: Resolved (решён), Owned (есть владелец), Accepted (принят), Mitigated (снижен). "
     "Ниже — высокие риски из реестров команд (вероятность × влияние ≥ 15) с кнопкой «↗ В ROAM» для переноса на уровень стрима.",
     ["Сводка по категориям ROAM над таблицей.", "Риск можно привязать к команде или ко всему стриму; указать владельца, срок и решение."],
     ["s_roam.jpg"], ["Риски ROAM с переносом высоких рисков команд"]),
    ("Scrum of Scrums",
     "Синхронизация команд стрима. «Сформировать повестку» создаёт строку на каждую команду на выбранную дату: прогресс текущего спринта, "
     "открытые импедименты и заблокированные зависимости команды, от кого нужна помощь (по межкомандным зависимостям).",
     ["Встречи выбираются по дате; строки повестки дополняются на встрече (сделано, планы, блокеры).",
      "Ниже — открытые межкомандные зависимости стрима."],
     ["s_sos.jpg"], ["Scrum of Scrums: повестка из данных команд"]),
    ("System Demo",
     "Совместный показ стрима в конце итерации. «Сформировать из фич итерации» добавляет в повестку фичи, истории которых готовы в выбранной "
     "итерации, с командами-участниками; затем фиксируются отзывы стейкхолдеров и итог «Принято / Доработать».",
     ["Сводка: показов, показано, принято и доля принятых.", "Повторное формирование не дублирует уже добавленные фичи."],
     ["s_demo.jpg"], ["System Demo: повестка, итоги и отзывы"]),
]
INTRO = ("Стримы. Начиная с версии 2.0 приложение работает с несколькими командами: команды объединяются в стримы (потоки ценности), "
         "у каждой команды свои данные, а на уровне стрима ведутся фичи, PI-планирование, межкомандные зависимости, риски ROAM, "
         "Scrum of Scrums и System Demo.")


def heading_idx(ps, prefix, style=None):
    return next((k for k, p in enumerate(ps) if p.style.name.startswith("Heading") and (style is None or p.style.name == style)
                 and p.text.strip().startswith(prefix)), None)


class Ins:
    """Последовательная вставка абзацев перед якорем (или в конец документа)."""
    def __init__(self, d, anchor): self.d, self.anchor = d, anchor

    def _put(self, p):
        if self.anchor is not None: self.anchor._p.addprevious(p._p)
        return p

    def par(self, text, style="Normal"):
        return self._put(self.d.add_paragraph(text, style=style))

    def pic(self, img, caption, width):
        if not (SH / img).exists(): print("   ! нет снимка", img); return
        p = self.d.add_paragraph(); p.alignment = 1; p.add_run().add_picture(str(SH / img), width=Inches(width)); self._put(p)
        c = self.d.add_paragraph(caption); c.alignment = 1
        for r in c.runs: r.italic = True
        self._put(c)


def section(ins, level_style, title, text, bullets, imgs, caps, width, with_pics, bullet_style="List Bullet"):
    ins.par(title, level_style); ins.par(text)
    for b in bullets: ins.par(b, bullet_style)
    if with_pics:
        for im, cp in zip(imgs, caps): ins.pic(im, cp, width)


def user_guide(d):
    ps = d.paragraphs
    if heading_idx(ps, "20. Стримы и команды"): return 0
    anchor = next((p for p in ps if p.style.name == "Heading 1" and p.text.startswith("Ширина столбцов")), None)
    ins = Ins(d, anchor)
    t, text, bullets, imgs, caps = ORG
    ins.par("20. Стримы и команды", "Heading 1"); ins.par(text)
    ins.par("20.1 Переключатель «Стрим › Команда»", "Heading 2")
    for b in bullets[:1]: ins.par(b, "List Bullet")
    ins.pic(imgs[1], caps[1], 6.3)
    ins.par("20.2 Раздел «Стримы и команды»", "Heading 2")
    for b in bullets[1:5]: ins.par(b, "List Bullet")
    ins.pic(imgs[0], caps[0], 6.3)
    ins.par("20.3 Роли в стримах и командах", "Heading 2")
    for b in [bullets[5], "Глобальный администратор видит и меняет всё; управлять стримами, командами и назначениями может только он.",
              "Роль в стриме даёт право менять данные стрима (фичи, PI, ROAM, Scrum of Scrums, System Demo); участник отдельной команды видит свой стрим только для просмотра.",
              "Права в интерфейсе следуют роли в текущей команде: при переключении команды режим «Просмотр/Редактор» меняется автоматически."]:
        ins.par(b, "List Bullet")
    ins.par("20.4 Локальный и серверный режим", "Heading 2")
    for b in ["Без сервера стримы, команды и данные стрима хранятся в этом браузере; роли не применяются.",
              "С сервером структура приходит из GET /api/org, данные команды — из её схемы БД (заголовок X-Team-Id), доска — своя комната у каждой команды."]:
        ins.par(b, "List Bullet")
    ins.par("21. Уровень стрима", "Heading 1")
    ins.par("Блок меню «Стрим» работает с текущим стримом (стрим текущей команды); стрим можно сменить списком в верхней части раздела.")
    for n, (title, text, bl, im, cp) in enumerate(STREAM, 1):
        section(ins, "Heading 2", f"21.{n} {title}", text, bl, im, cp, 6.3, True)
    ins.par("21.7 Типовой цикл PI", "Heading 2")
    for s in ["Перед PI: заполните «Фичи стрима», оцените WSJF и назначьте фичам PI.",
              "PI-планирование: разбейте фичи на истории в бэклогах команд («＋ История»), зафиксируйте цели команд и межкомандные зависимости; проверьте конфликты на программной доске.",
              "В ходе PI: проводите Scrum of Scrums, ведите риски ROAM, следите за burn-up PI на обзоре стрима.",
              "В конце итерации: System Demo по фичам с готовыми историями; в конце PI — факт BV целей и предсказуемость."]:
        ins.par(s, "List Number")
    # оглавление + «Начало работы»
    ps = d.paragraphs
    toc = next((k for k, p in enumerate(ps) if p.text.startswith("17. Справочник элементов")), None)
    if toc is not None:
        a = ps[toc + 1]
        for line in ("18. Практики проектного управления", "19. Здоровье и производительность команды", "20. Стримы и команды", "21. Уровень стрима"):
            if not any(p.text == line for p in ps[toc:toc + 6]): a.insert_paragraph_before(line, style=ps[toc].style)
    i = heading_idx(d.paragraphs, "1. Начало работы", "Heading 1")
    if i is not None: d.paragraphs[i + 1].insert_paragraph_before(INTRO + " Подробно — главы 20 и 21.")
    return 1


def description(d, pics):
    ps = d.paragraphs
    if heading_idx(ps, MARK, "Heading 2"): return 0
    w = 6.61
    ob = heading_idx(ps, "Обзор", "Heading 1")
    first = ps[ob + 1]                                  # первый раздел «Обзора» — перед ним раздел «Стримы и команды»
    t, text, bullets, imgs, caps = ORG
    section(Ins(d, first), "Heading 2", t, text, bullets, imgs, caps, w, pics)
    ps = d.paragraphs
    pl = next(p for p in ps if p.style.name == "Heading 1" and p.text.strip() == "Планирование")
    ins = Ins(d, pl)
    ins.par("Стрим", "Heading 1")
    for title, text, bl, im, cp in STREAM: section(ins, "Heading 2", title, text, bl, im, cp, w, pics)
    i = next((k for k, p in enumerate(d.paragraphs) if p.style.name == "Heading 1" and p.text.startswith("Кросс-функциональная команда")), None)
    if i is not None:
        j = next(k for k in range(i + 1, len(d.paragraphs)) if d.paragraphs[k].style.name.startswith("Heading"))
        d.paragraphs[j].insert_paragraph_before(INTRO)
    return 1


DEP_CH = ("14. Стримы и команды: особенности развёртывания", [
    "Порт и изоляция: веб-интерфейс публикуется на 8090, compose-проект — upravlenie-streamami, БД — streams и streams_audit; приложение можно запускать рядом с «Управлением кросс-функциональной командой» (8080).",
    "Миграции при старте: сначала таблицы команд из db/migration (V1…V12) в схеме public, затем платформенные таблицы из db/platform (своя история platform_schema_history: V1 — org_streams, org_teams, org_memberships; V2 — org_stream_docs), затем все схемы команд team_<id>.",
    "Первая команда (t1) использует схему public, каждая новая команда получает схему team_<id>, которую создаёт и мигрирует сам сервис. Пользователю БД нужно право CREATE на базу (в docker-compose он владелец базы).",
    "Проверка: GET /api/org (под администратором) возвращает стримы и команды; создайте тестовую команду в разделе «Стримы и команды» и убедитесь, что GET /api/state с заголовком X-Team-Id возвращает её данные.",
    "Резервное копирование: pg_dump всей базы streams включает все схемы команд; копия одной команды — pg_dump -n team_<id> streams. Удаление команды удаляет её схему безвозвратно — делайте копию заранее.",
    "Обновление версии: схемы всех команд мигрируются автоматически при старте сервиса; исправления базового продукта переносятся через git upstream (git fetch upstream && git merge upstream/main).",
    "Nginx пропускает заголовок X-Team-Id и параметр team у /ws/board без дополнительной настройки."])
ADM_CH = ("13. Стримы, команды и роли", [
    "Модель: Портфель → Стрим → Команда. Таблицы платформы (схема public): org_streams, org_teams (schema_name — схема данных команды, пусто у первой), org_memberships (пользователь, область STREAM/TEAM, роль VIEWER/EDITOR), org_stream_docs (документ стрима с версией).",
    "Данные команды: каждая команда — отдельная схема с полным набором таблиц приложения (первая — public). Запросы /api/state, /api/tasks, /analytics/* выполняются в схеме команды из заголовка X-Team-Id.",
    "Эффективная роль на команде: ADMIN — всё; пользователь без назначений — его общая роль на всех командах; иначе — максимум из роли в команде и роли в её стриме; нет назначения — нет доступа (403).",
    "Роль в стриме (данные стрима): ADMIN; без назначений — общая роль; назначение на стрим — его роль; участник команды стрима — только просмотр.",
    "Управление: раздел «Стримы и команды» (блок «Роли в стримах и командах») или API: GET /api/org; POST/PUT/DELETE /api/org/streams[/{id}], /api/org/teams[/{id}] (template: empty | copy, source); GET/POST/DELETE /api/org/memberships[/{pk}]; GET/PUT /api/org/streams/{id}/doc ({data, baseVersion}, 409 при конфликте версий).",
    "Ограничения: первую команду удалить нельзя; стрим удаляется только пустым; удаление команды удаляет её схему и назначения на неё.",
    "Совместная доска: отдельная комната у каждой команды (/ws/board?team=<id>), подключение проверяет доступ к команде.",
    "Неполадки: 404 team_not_found — неверный X-Team-Id (проверьте GET /api/org); 403 — нет роли на команде/стриме; 409 — данные команды или стрима изменил другой пользователь (клиент загружает свежую версию)."])


def chapter(d, title, bullets):
    if heading_idx(d.paragraphs, title.split(".")[0] + ". "+ title.split(". ", 1)[1]): return 0
    ins = Ins(d, None); ins.par(title, "Heading 1")
    for b in bullets: ins.par(b, "List Bullet")
    return 1


for fn in (UG, ILL, DESC, DEP, ADM):
    d = docx.Document(fn)
    n = sum(fix_par(p) for p in all_pars(d))
    added = {UG: lambda: user_guide(d), ILL: lambda: description(d, True), DESC: lambda: description(d, False),
             DEP: lambda: chapter(d, *DEP_CH), ADM: lambda: chapter(d, *ADM_CH)}[fn]()
    d.save(fn)
    print(f"{fn[:62]:62s} замен: {n:3d}  новые главы: {'да' if added else 'уже есть'}")
