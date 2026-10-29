# -*- coding: utf-8 -*-
"""v2.4.0: «Обзор стрима» → «Дашборд стрима» в документах.
Переписывает раздел (руководство пользователя 21.1, описание и иллюстрированное описание) — текст, пункты и снимки
build/shots/s_stream, s_charts, s_matrix, s_cards, s_flow, s_quality (capture_stream_shots.py); заменяет упоминания
«обзор стрима» и версию во всех .docx. Идемпотентно. Слайды — release_v23_slides.py.
Запуск: python scripts/deliverables/release_v24_dashboard.py"""
import glob, os, pathlib, re, docx
from docx.shared import Inches

os.chdir(pathlib.Path(__file__).resolve().parents[2])
SH = pathlib.Path("build/shots")
P = "Управление_стримами_и_командами"
OLDV, NEWV = "2.3.0", "2.4.0"
TEXT = ("Все показатели стрима и каждой команды на одной странице. Показатели команды считаются теми же формулами, что и её "
        "дашборд (velocity, health, риск срыва спринта, say/do, cycle time, DORA, Team Health Check и др.); показатели стрима — "
        "сводка по командам и данным стрима (предсказуемость PI, фичи, burn-up, System Demo, зависимости, ROAM).")
BULLETS = [
    "«Требует внимания»: высокий риск срыва спринта, перегруз ёмкости, критичные баги, импедименты старше 7 дней, красные зоны Health Check, "
    "конфликты межкомандных зависимостей, нерешённые высокие риски ROAM, фичи PI без историй, предсказуемость ниже 80%. Клик по сигналу открывает команду или раздел с причиной.",
    "«Сводка»: KPI стрима, burn-up PI, velocity команд по итерациям и матрица команд.",
    "«Команды»: тепловая матрица (цвет ячейки — светофор показателя) или карточки с мини-трендом velocity. Клик по команде открывает боковую панель "
    "с причинами риска и всеми показателями по группам; «Открыть команду» переключает на её данные.",
    "«Поток»: velocity, cycle time (медиана), WIP и старейшая задача в работе, say/do, загрузка, распределение работы (фичи и истории, техдолг, дефекты).",
    "«Качество»: открытые баги по серьёзности, DRE, DoD, DORA (деплои в неделю, lead time, change failure rate, MTTR), техдолг.",
    "«Люди»: численность, Team Health Check, настроение, одиночные носители навыков (SPOF), импедименты и их возраст.",
    "«Тренды»: срезы показателей по итерациям PI — автоматически при первом открытии дашборда в итерации и кнопкой «📌 Зафиксировать срез»; "
    "графики velocity, health, предсказуемости, say/do, багов и блокеров. Срезы выгружаются в Excel (лист «Срезы стрима»).",
]
PICS = [("s_stream.jpg", "Дашборд стрима: сигналы «Требует внимания» и KPI"), ("s_charts.jpg", "Сводка: burn-up PI, velocity команд и матрица команд"),
        ("s_matrix.jpg", "Вкладка «Команды»: тепловая матрица"), ("s_cards.jpg", "Карточки команд и боковая панель команды"),
        ("s_flow.jpg", "Вкладка «Поток»"), ("s_quality.jpg", "Вкладка «Качество»")]
MARK = PICS[3][1]


def rewrite(d, title_re, new_title, width, pics):
    ps = d.paragraphs
    i = next((k for k, p in enumerate(ps) if p.style.name.startswith("Heading") and re.match(title_re, p.text.strip())), None)
    if i is None: return "нет раздела"
    if any(p.text == MARK for p in ps): return "уже обновлён"
    h = ps[i]; lvl = h.style.name
    j = next((k for k in range(i + 1, len(ps)) if ps[k].style.name.startswith("Heading")), None)
    for p in ps[i + 1:j if j is not None else len(ps)]: p._p.getparent().remove(p._p)
    num = re.match(r"(\d+(?:\.\d+)?\s)", h.text)
    for r in h.runs[1:]: r.text = ""
    h.runs[0].text = (num.group(1) if num else "") + new_title
    anchor = d.paragraphs[i + 1] if j is not None else None

    def put(p):
        if anchor is not None: anchor._p.addprevious(p._p)
        return p
    put(d.add_paragraph(TEXT))
    for b in BULLETS: put(d.add_paragraph(b, style="List Bullet"))
    if pics:
        for img, cap in PICS:
            p = d.add_paragraph(); p.alignment = 1; p.add_run().add_picture(str(SH / img), width=Inches(width)); put(p)
            c = d.add_paragraph(cap); c.alignment = 1
            for r in c.runs: r.italic = True
            put(c)
    return "обновлён (" + lvl + ")"


def fix(d):
    n = 0
    def all_pars():
        yield from d.paragraphs
        for t in d.tables:
            for row in t.rows:
                for c in row.cells: yield from c.paragraphs
        for s in d.sections:
            for hf in (s.header, s.footer): yield from hf.paragraphs
    for p in all_pars():
        for a, b in (("обзоре стрима", "дашборде стрима"), ("«Обзор стрима»", "«Дашборд стрима»"), ("Обзор стрима", "Дашборд стрима"), (OLDV, NEWV)):
            if a in p.text:
                for r in p.runs: r.text = r.text.replace(a, b)
                if a in p.text:
                    full = p.text.replace(a, b); p.runs[0].text = full
                    for r in p.runs[1:]: r.text = ""
                n += 1
    return n


for fn, title_re, width, pics in ((f"Руководство_пользователя_{P}.docx", r"21\.1 Обзор стрима", 6.3, True),
                                  (f"Иллюстрированное_описание_{P}.docx", r"Обзор стрима$", 6.61, True),
                                  (f"Описание_{P}.docx", r"Обзор стрима$", 6.61, False)):
    d = docx.Document(fn); res = rewrite(d, title_re, "Дашборд стрима", width, pics); d.save(fn); print(f"{fn[:58]:58s} раздел: {res}")
for fn in sorted(glob.glob(f"*_{P}.docx")):
    d = docx.Document(fn); n = fix(d); d.save(fn); print(f"{fn[:58]:58s} замен: {n}")
