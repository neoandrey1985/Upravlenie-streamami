# -*- coding: utf-8 -*-
"""Презентация_приложения.pptx для «Управления стримами и командами» v2.3.0.
Титульный и итоговый слайды — под новый продукт; после слайда «Кросс-функциональная команда» — слайд-концепция
«Стрим» и 9 слайдов уровня стрима со снимками capture_stream_shots.py (build/shots/s_*.jpg). Оформление клонируется
с существующих слайдов. Идемпотентно (слайды узнаются по заголовку).
Запуск: python scripts/deliverables/release_v23_slides.py"""
import copy, os, pathlib
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

os.chdir(pathlib.Path(__file__).resolve().parents[2])
DECK, SH = "Презентация_приложения.pptx", pathlib.Path("build/shots")
CONCEPT = ("🛰️", "Стрим — несколько команд, один поток ценности",
           ["Стрим (поток ценности, Agile Release Train) объединяет 5 и более кросс-функциональных команд, которые вместе поставляют ценность заказчику.",
            "У каждой команды свои бэклог, спринты, доска и метрики, но планируются и синхронизируются они вместе: общий бэклог фич, PI-планирование, межкомандные зависимости и риски."],
           "Приложение ведёт каждую команду отдельно (своя схема данных) и собирает их в стрим: обзор стрима, фичи с WSJF, PI-доска, ROAM, Scrum of Scrums и System Demo.")
SLIDES = [
 ("🏢", "Стримы и команды", "Оргструктура «Стрим → Команда».", "s_org.jpg", [
   "•  Переключатель «Стрим › Команда» в шапке — данные каждой команды отдельно.",
   "•  Команда создаётся пустой или копией; импорт команды из JSON.",
   "•  Роли на стрим и на команду; без назначений — общая роль.",
   "•  На сервере — своя схема БД у каждой команды."]),
 ("🛰️", "Обзор стрима", "Все команды стрима на одном экране.", "s_stream.jpg", [
   "•  KPI стрима: люди, velocity, health, предсказуемость PI, баги, блокеры.",
   "•  Таблица команд: спринт, риск срыва, health, say/do — с RAG.",
   "•  Те же формулы, что и на дашборде команды.",
   "•  «Открыть» — переход в данные команды."]),
 ("📈", "Графики стрима", "Burn-up PI и velocity по итерациям.", "s_charts.jpg", [
   "•  Burn-up: объём историй фич PI, сделанное и идеальный темп.",
   "•  Velocity каждой команды по итерациям PI.",
   "•  Фичи текущего PI по WSJF с прогрессом."]),
 ("🎁", "Фичи стрима", "Бэклог фич с WSJF — таблицей или канбаном.", "s_kanban.jpg", [
   "•  WSJF = (BV + TC + RR) / размер; статус, PI, владелец.",
   "•  Канбан: Воронка → Анализ → Бэклог → В работе → Готово.",
   "•  «＋ История» — история фичи сразу в бэклоге нужной команды.",
   "•  Прогресс фичи — по статусам её историй в командах."]),
 ("🗓️", "PI стрима: программная доска", "Команды × итерации инкремента.", "s_pi.jpg", [
   "•  Доска строится из историй фич автоматически.",
   "•  Готовые фичи — зелёные, зависимости — 🔗 на доске.",
   "•  Параметры PI: старт, число итераций, несколько PI."]),
 ("🔗", "Цели PI и межкомандные зависимости", "Предсказуемость и контроль сроков.", "s_deps.jpg", [
   "•  Цели команд с плановой и фактической бизнес-ценностью.",
   "•  Предсказуемость PI = факт / план по обязательным целям.",
   "•  Зависимости: кто ждёт, кто делает, к какой итерации.",
   "•  Конфликт сроков и блокировки подсвечиваются красным."]),
 ("🛡️", "Риски ROAM", "Риски уровня стрима.", "s_roam.jpg", [
   "•  Resolved, Owned, Accepted, Mitigated — сводка по категориям.",
   "•  Высокие риски команд переносятся одной кнопкой.",
   "•  Владелец, срок и решение по каждому риску."]),
 ("🔄", "Scrum of Scrums", "Синхронизация команд стрима.", "s_sos.jpg", [
   "•  Повестка формируется из данных команд автоматически.",
   "•  Прогресс спринта, импедименты, блокеры, кому нужна помощь.",
   "•  Открытые межкомандные зависимости под рукой."]),
 ("🎬", "System Demo", "Совместный показ стрима в конце итерации.", "s_demo.jpg", [
   "•  Повестка — фичи с готовыми в итерации историями.",
   "•  Отзывы стейкхолдеров и итог «Принято / Доработать».",
   "•  Доля принятых фич по PI."]),
]
TITLE = ("Управление стримами и командами",
         "Веб-приложение для agile-планирования стримов из нескольких кросс-функциональных команд: данные команд, уровень стрима, метрики и диаграммы в одном месте.")
SUMMARY = ("Единый инструмент для стримов из нескольких команд: от бэклога каждой команды до PI-планирования стрима — с общей БД и понятной архитектурой.",
           ["•  Стримы и команды: своя схема данных, роли на стрим и команду",
            "•  Уровень стрима: фичи, PI-доска, ROAM, Scrum of Scrums, System Demo"],
           "github.com/neoandrey1985/Upravlenie-streamami")


def texts(s): return [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
def title_of(s):
    t = texts(s); return t[1].text_frame.text.strip() if len(t) > 1 else (t[0].text_frame.text.strip() if t else "")


def set_text(shape, text):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p'))
    for p in ps[1:]: tb.remove(p)
    runs = ps[0].findall(qn('a:r'))
    for r in runs[1:]: ps[0].remove(r)
    runs[0].find(qn('a:t')).text = text


def set_paras(shape, lines):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p')); tpl = copy.deepcopy(ps[0])
    for p in ps: tb.remove(p)
    for ln in lines:
        p = copy.deepcopy(tpl); runs = p.findall(qn('a:r'))
        for r in runs[1:]: p.remove(r)
        runs[0].find(qn('a:t')).text = ln; tb.append(p)


def parts(s):
    t = texts(s); icon = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.has_text_frame)
    body = next(sh for sh in t if sh.text_frame.text.strip().startswith("•"))
    rest = [sh for sh in t if sh._element is not icon._element and sh._element is not body._element]
    pic = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    return icon, rest[0], rest[1], body, pic


def clone_after(prs, src, pos, shot=None):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None: new._element.cSld.insert(0, copy.deepcopy(bg))
    pic = None
    for sh in src.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE: pic = sh; continue
        new.shapes._spTree.append(copy.deepcopy(sh._element))
    if pic is not None and shot: new.shapes.add_picture(str(SH / shot), pic.left, pic.top, pic.width, pic.height)
    lst = prs.slides._sldIdLst; el = lst[-1]; lst.remove(el); lst.insert(pos, el)
    return new


def set_picture(slide, pic, path):
    slide.part.related_part(pic._element.find('.//' + qn('a:blip')).get(qn('r:embed')))._blob = pathlib.Path(path).read_bytes()


prs = Presentation(DECK)
sl = list(prs.slides)
# титульный слайд
t1 = texts(sl[0]); set_text(t1[0], TITLE[0]); set_text(t1[1], TITLE[1])
# концепция «Стрим» после слайда «Кросс-функциональная команда» (клон его оформления)
ci = next(i for i, s in enumerate(sl) if title_of(s) == "Кросс-функциональная команда")
cs = next((s for s in prs.slides if title_of(s) == CONCEPT[1]), None)
if cs is None: cs = clone_after(prs, sl[ci], ci + 1)
ct = texts(cs); set_text(ct[0], CONCEPT[0]); set_text(ct[1], CONCEPT[1]); set_paras(ct[2], CONCEPT[2]); set_text(ct[3], CONCEPT[3])
# слайды уровня стрима (оформление — слайд «Дашборд»)
src = next(s for s in prs.slides if title_of(s) == "Дашборд")
pos0 = list(prs.slides).index(cs) + 1
for k, (ic, title, sub, shot, bullets) in enumerate(SLIDES):
    s = next((x for x in prs.slides if title_of(x) == title), None)
    if s is None: s = clone_after(prs, src, pos0 + k, shot)
    icon, tsh, ssh, bsh, psh = parts(s)
    set_text(icon, ic); set_text(tsh, title); set_text(ssh, sub); set_paras(bsh, bullets); set_picture(s, psh, SH / shot)
# итог
last = list(prs.slides)[-1]; lt = texts(last)
set_text(lt[1], SUMMARY[0])
body = next(sh for sh in lt if sh.text_frame.text.strip().startswith("•"))
set_paras(body, SUMMARY[1] + [
    "•  Для каждой команды — Kanban, RICE, церемонии, доска и 50+ разделов",
    "•  Совместная доска в реальном времени — своя у каждой команды",
    "•  Роли в PostgreSQL, проверка прав в Java и Python",
    "•  Java + Python + PostgreSQL, CI и docker-compose"])
link = next((sh for sh in lt if "github.com" in sh.text_frame.text), None)
if link is not None: set_text(link, SUMMARY[2])
prs.save(DECK)
for i, s in enumerate(prs.slides, 1):
    if i <= 14 or i == len(prs.slides): print(i, title_of(s))
