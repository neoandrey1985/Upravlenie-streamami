# -*- coding: utf-8 -*-
"""v2.5.0: раздел «Зрелость стрима» в документах. Руководство пользователя — 21.8, описания — раздел в группе «Стрим»
(снимки build/shots/s_maturity, s_mat_practices, s_mat_teams), руководство администратора — метод голосования,
версия во всех .docx. Идемпотентно.  Запуск: python scripts/deliverables/release_v25_maturity.py"""
import glob, os, pathlib, docx
from docx.shared import Inches

os.chdir(pathlib.Path(__file__).resolve().parents[2])
SH = pathlib.Path("build/shots")
P = "Управление_стримами_и_командами"
OLDV, NEWV = "2.4.0", "2.5.0"
TITLE = "Зрелость стрима"
TEXT = ("Agile-зрелость стрима в трёх слоях: зрелость команд (их самооценки по Scrum, Kanban и SAFe в разрезе стрима), практики стрима "
        "(10 измерений с рубрикой из 5 уровней — от «Начального» до «Оптимизирующего») и объективные индикаторы из данных приложения, "
        "которые проверяют самооценку. Оценка проводится раз в PI, на Inspect & Adapt.")
BULLETS = [
    "Измерения практик стрима: PI-планирование, синхронизация каденции, бэклог фич, межкомандные зависимости, System Demo, Inspect & Adapt, "
    "роли стрима, непрерывная поставка, технические практики, метрики и прозрачность. Ненужные можно отключить и добавить свои («Настроить измерения»).",
    "Рубрика: у каждого измерения описан каждый из 5 уровней (раскройте измерение в таблице); текущий уровень выделен.",
    "Голоса команд: каждая команда ставит уровень 1–5 по рубрике; участник команды голосует за свою команду (с сервером — даже без роли в стриме). "
    "Показываются медиана и разброс голосов; разброс от 2 уровней подсвечивается.",
    "Итог RTE: пользователь с правом изменения данных стрима фиксирует итоговый уровень и обоснование; «Принять медианы» ставит итог по медиане "
    "там, где он не задан. Пока итога нет, используется медиана голосов.",
    "Индикаторы из данных: предсказуемость PI, доля команд в каденции PI, фичи с WSJF и историями, конфликты зависимостей, регулярность System Demo, "
    "разбор рисков ROAM и выполнение улучшений, уровни DORA, DoD команд, срезы показателей и встречи Scrum of Scrums. Если итог расходится с индикатором "
    "на 2 уровня и более, сверка помечает «возможна переоценка» или «возможна недооценка».",
    "Итоговый уровень стрима — медиана по измерениям, «слабое звено» — самое низкое измерение; радар «итог ↔ индикаторы» и динамика по PI.",
    "Зрелость команд: тепловая карта команд по измерениям выбранного фреймворка (Scrum, Kanban, SAFe) со средним, минимумом и разбросом.",
    "Бэклог улучшений: действия по измерениям с владельцем, PI и статусом; «Предложить из слабых измерений» создаёт действия для 2–3 самых "
    "слабых измерений с формулировкой следующего уровня рубрики.",
    "На дашборде стрима — плитка «Зрелость стрима» и сигналы о расхождении оценки с фактами и о непроведённой оценке в конце PI; "
    "в Excel — листы «Зрелость стрима» и «Улучшения стрима».",
]
PICS = [("s_maturity.jpg", "Зрелость стрима: итоговая оценка и индикаторы"),
        ("s_mat_practices.jpg", "Практики стрима: голоса, итог RTE, индикаторы и сверка"),
        ("s_mat_teams.jpg", "Голоса команд и зрелость команд")]
ADM = ("Голос команды в оценке зрелости стрима: POST /api/org/streams/{id}/vote с телом {team, pi, votes: {измерение: уровень 1–5, 0 — снять}}. "
       "Нужна роль EDITOR на этой команде; сервер сливает голос в документ стрима под блокировкой строки и повышает его версию.")


def insert(d, anchor, title_text, style, width, pics):
    def put(p):
        if anchor is not None: anchor._p.addprevious(p._p)
        return p
    put(d.add_paragraph(title_text, style=style)); put(d.add_paragraph(TEXT))
    for b in BULLETS: put(d.add_paragraph(b, style="List Bullet"))
    if pics:
        for img, cap in PICS:
            p = d.add_paragraph(); p.alignment = 1; p.add_run().add_picture(str(SH / img), width=Inches(width)); put(p)
            c = d.add_paragraph(cap); c.alignment = 1
            for r in c.runs: r.italic = True
            put(c)


def has_heading(d, text):
    return any(p.style.name.startswith("Heading") and p.text.strip().endswith(text) for p in d.paragraphs)


def version(d):
    n = 0
    pars = list(d.paragraphs) + [p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs] \
        + [p for s in d.sections for hf in (s.header, s.footer) for p in hf.paragraphs]
    for p in pars:
        if OLDV in p.text:
            for r in p.runs: r.text = r.text.replace(OLDV, NEWV)
            n += 1
    return n


# руководство пользователя: 21.8 — в конце главы 21
fn = f"Руководство_пользователя_{P}.docx"; d = docx.Document(fn)
if not has_heading(d, TITLE):
    ps = d.paragraphs; i = next(k for k, p in enumerate(ps) if p.style.name == "Heading 1" and p.text.startswith("21. "))
    anchor = next((p for p in ps[i + 1:] if p.style.name == "Heading 1"), None)
    insert(d, anchor, "21.8 " + TITLE, "Heading 2", 6.3, True); st = "добавлен"
else: st = "уже есть"
d.save(fn); print(fn[:58], "раздел:", st)
# описания: в конце группы «Стрим» (перед «Планированием»)
for fn, pics in ((f"Иллюстрированное_описание_{P}.docx", True), (f"Описание_{P}.docx", False)):
    d = docx.Document(fn)
    if not has_heading(d, TITLE):
        anchor = next(p for p in d.paragraphs if p.style.name == "Heading 1" and p.text.strip() == "Планирование")
        insert(d, anchor, TITLE, "Heading 2", 6.61, pics); st = "добавлен"
    else: st = "уже есть"
    d.save(fn); print(fn[:58], "раздел:", st)
# руководство администратора: метод голосования в главе 13
fn = f"Руководство_администратора_{P}.docx"; d = docx.Document(fn)
if not any(p.text.startswith("Голос команды в оценке зрелости") for p in d.paragraphs):
    ref = next(p for p in d.paragraphs if p.text.startswith("Управление: раздел «Стримы и команды»"))
    new = ref.insert_paragraph_before(ADM, style=ref.style)
    ref._p.addnext(new._p)   # после пункта «Управление…»
    st = "добавлен"
else: st = "уже есть"
d.save(fn); print(fn[:58], "API голосования:", st)
for fn in sorted(glob.glob(f"*_{P}.docx")):
    d = docx.Document(fn); n = version(d); d.save(fn); print(f"{fn[:58]:58s} версия: {n}")
